"""Read a local Source Data workbook without changing it."""
from pathlib import Path
import hashlib
import json
import numbers
import numpy as np
import pandas as pd
from openpyxl import load_workbook

REQUIRED = ("Table S4", "Table S5", "Retrospective patient metrics", "Patient-level OAR averages",
            "Prospective patient metrics", "Prospective timing", "External-centre summary", "Fig S4 GTVn comparison")


def read_source(path):
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Source Data not found: {path}. Check the explicitly supplied --source-data path. Omit --source-data to plot using the bundled data.")
    result = {}
    with path.open("rb") as handle:
        book = load_workbook(handle, read_only=True, data_only=True)
        try:
            missing = set(REQUIRED)-set(book.sheetnames)
            if missing:
                raise ValueError(f"Missing worksheets: {sorted(missing)}")
            for name in REQUIRED:
                rows = list(book[name].values)
                width = 7 if name == "Fig S4 GTVn comparison" else len(rows[0])
                header = rows[0][:width]
                frame = pd.DataFrame([r[:width] for r in rows[1:]], columns=header)
                frame["source_row"] = np.arange(2, len(frame)+2)
                if frame[list(header)].isna().all(axis=1).any():
                    raise ValueError(f"Unexpected empty data row in {name}")
                result[name] = frame
            # The strategy summary occupies a separate block on the same sheet.
            rows = list(book["Fig S4 GTVn comparison"].values)
            result["GTVn stored summary"] = pd.DataFrame([r[8:14] for r in rows[1:] if len(r)>8 and r[8] is not None], columns=rows[0][8:14])
        finally:
            book.close()
    result["metadata"] = {"source_filename": path.name, "source_sha256": hashlib.sha256(path.read_bytes()).hexdigest()}
    return result


def require_unique(frame, keys, context):
    if frame[keys].isna().any().any() or (frame[keys].astype(str).apply(lambda c: c.str.strip().eq("")).any().any()):
        raise ValueError(f"Missing identifier in {context}")
    duplicated = frame.duplicated(keys, keep=False)
    if duplicated.any():
        raise ValueError(f"Duplicate key in {context}: {frame.loc[duplicated, keys].head().to_dict('records')}")


def require_numeric(frame, columns, context):
    for column in columns:
        # Reject numeric-looking strings, booleans, dates and empty cells.
        if not frame[column].map(lambda x: isinstance(x, numbers.Real) and not isinstance(x, (bool, np.bool_)) and np.isfinite(x)).all():
            raise ValueError(f"Non-numeric or missing value in {context}/{column}")


def write_csv(frame, path):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(frame).to_csv(path, index=False, encoding="utf-8", float_format="%.17g")


def write_json(value, path):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False), encoding="utf-8")
